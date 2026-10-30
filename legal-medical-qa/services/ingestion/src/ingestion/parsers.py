"""Format verification and parsers (doc 11.1 steps 3 and 4).

Magic bytes are checked against the declared type before any parser runs. Structured sources are
preferred (JATS, SPL, eCFR XML, DOCX, sanitized HTML); PDFs use the text layer, and a page with
no usable text layer is flagged ``extraction='none'`` so the pipeline quarantines it (OCR review)
rather than embedding nothing. Every parser is deterministic so text can be re-extracted.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import TYPE_CHECKING, Any, Final, cast

import pypdfium2 as pdfium
from lxml import etree, html

from ingestion.headings import classify_line
from ingestion.model import Block, PageText, ParsedDocument

if TYPE_CHECKING:
    from lxml.html import HtmlElement

MAX_PAGES: Final = 5000
MAX_BYTES: Final = 200 * 1024 * 1024
MIN_USABLE_CHARS: Final = 25
_PAGE_BREAK: Final = re.compile(r"page-break-(?:before|after)\s*:\s*always", re.IGNORECASE)
_BLOCK_TAGS: Final = {
    "p",
    "div",
    "li",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "blockquote",
    "pre",
    "center",
}
_SPL_CODES: Final = {
    "34067-9": "Indications and usage",
    "34068-7": "Dosage and administration",
    "34070-3": "Contraindications",
    "34071-1": "Warnings",
    "43685-7": "Warnings and precautions",
    "34084-4": "Adverse reactions",
    "34073-7": "Drug interactions",
    "34090-1": "Clinical pharmacology",
}


class UnsupportedFormatError(ValueError):
    """The bytes do not match the declared type, or the type is not supported."""


def detect_format(data: bytes, declared: str) -> str:
    """Verify magic bytes against the declared type and return the canonical format name."""
    if len(data) > MAX_BYTES:
        msg = "file exceeds the size limit"
        raise UnsupportedFormatError(msg)
    head = data[:2048].lstrip()
    declared = declared.lower().lstrip(".")
    if data[:5] == b"%PDF-":
        found = "pdf"
    elif data[:4] == b"PK\x03\x04" and zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            found = "docx" if "word/document.xml" in z.namelist() else "zip"
    elif head.startswith((b"<?xml", b"<article", b"<document", b"<ECFR", b"<DIV")) or (
        head[:9].lower() == b"<!doctype" and b"html" not in head[:30].lower()
    ):
        found = "xml"  # JATS opens with <!DOCTYPE article PUBLIC ...>; the DTD is never loaded
    elif b"<html" in data[:4096].lower() or head.lower().startswith(b"<!doctype html"):
        found = "html"  # includes the SEC EDGAR <DOCUMENT><TEXT><HTML> submission wrapper
    elif b"\x00" not in data[:4096]:
        found = "txt"
    else:
        found = "binary"
    allowed = {
        "pdf": {"pdf"},
        "docx": {"docx"},
        "xml": {"xml", "jats", "spl", "ecfr"},
        "html": {"html", "htm"},
        "txt": {"txt", "text", "md"},
    }
    if found not in allowed or declared not in allowed[found]:
        msg = f"content looks like {found!r} but the declared type is {declared!r}"
        raise UnsupportedFormatError(msg)
    return found


# ---------------------------------------------------------------- shared builder
class _Builder:
    """Accumulates one page's text and the blocks inside it."""

    def __init__(self) -> None:
        self.pages: list[PageText] = []
        self.blocks: list[Block] = []
        self._text: list[str] = []
        self._len = 0
        self._page_blocks: list[
            tuple[int, int, str, int, int | None, str, str | None, str | None]
        ] = []

    def add(  # noqa: PLR0913 - a block genuinely has this many structural attributes
        self,
        text: str,
        kind: str = "paragraph",
        *,
        level: int = 0,
        table_id: int | None = None,
        section_path: str = "",
        clause_no: str | None = None,
        heading: str | None = None,
    ) -> None:
        text = text.strip()
        if not text:
            return
        start = self._len
        self._text.append(text)
        self._len += len(text)
        self._page_blocks.append(
            (start, self._len, kind, level, table_id, section_path, clause_no, heading)
        )
        self._text.append("\n")
        self._len += 1

    def end_page(self, *, extraction: str = "structured", force: bool = False) -> None:
        if not self._page_blocks and not force:
            return  # a page break with nothing before it does not create an empty page
        number = len(self.pages) + 1
        self.pages.append(PageText(number, "".join(self._text), extraction))
        for s, e, kind, level, table_id, path, clause, heading in self._page_blocks:
            self.blocks.append(Block(number, s, e, kind, level, table_id, path, clause, heading))
        self._text, self._len, self._page_blocks = [], 0, []

    def build(self, fmt: str, domain: str) -> ParsedDocument:
        if self._page_blocks:
            self.end_page()
        detected = any(b.kind == "heading" for b in self.blocks)
        return ParsedDocument(
            tuple(self.pages),
            tuple(self.blocks),
            fmt,
            "detected" if detected else "fallback",
            domain,
        )


def _line_blocks(builder: _Builder, lines: list[str], path_stack: list[str]) -> None:
    """Classify each line as heading or body; grouping into units is the chunker's job."""
    for line in lines:
        cls = classify_line(line)
        if cls is None:
            builder.add(line, "paragraph", section_path=" > ".join(path_stack))
            continue
        del path_stack[cls.level :]
        path_stack.append(cls.label)
        builder.add(
            line,
            "heading",
            level=cls.level,
            section_path=" > ".join(path_stack),
            clause_no=cls.number,
            heading=cls.title,
        )


# ---------------------------------------------------------------- PDF
def parse_pdf(data: bytes) -> ParsedDocument:
    """Use the text layer per page; a page without usable text is flagged for OCR review."""
    pdf = pdfium.PdfDocument(data)
    try:
        if len(pdf) > MAX_PAGES:
            msg = "page limit exceeded"
            raise UnsupportedFormatError(msg)
        builder, stack = _Builder(), []
        for index in range(len(pdf)):
            textpage = pdf[index].get_textpage()
            raw = textpage.get_text_range()
            textpage.close()
            lines = [ln.strip() for ln in raw.replace("\r", "\n").split("\n") if ln.strip()]
            usable = sum(len(ln) for ln in lines) >= MIN_USABLE_CHARS
            if usable:
                _line_blocks(builder, lines, stack)
            builder.end_page(extraction="text_layer" if usable else "none", force=True)
        return builder.build("pdf", "legal")
    finally:
        pdf.close()


# ---------------------------------------------------------------- HTML
def _text_of(el: HtmlElement) -> str:
    return " ".join(" ".join(el.itertext()).split())


def _walk_html(el: HtmlElement, builder: _Builder, stack: list[str], tables: list[int]) -> None:
    style = (el.get("style") or "").lower()
    forced = bool(_PAGE_BREAK.search(style))
    if forced and "before" in style:
        builder.end_page()
    if el.tag == "table":
        tables.append(len(tables))
        for row in el.iter("tr"):
            cells = [t for t in (_text_of(c) for c in row if c.tag in {"td", "th"}) if t]
            if cells:
                builder.add(
                    " | ".join(cells),
                    "table_row",
                    table_id=tables[-1],
                    section_path=" > ".join(stack),
                )
    else:
        has_block_children = any(
            isinstance(c.tag, str) and (c.tag in _BLOCK_TAGS or c.tag == "table") for c in el
        )
        if el.tag in _BLOCK_TAGS and not has_block_children:
            _line_blocks(builder, [_text_of(el)], stack)
        else:
            if el.tag in _BLOCK_TAGS and el.text and el.text.strip():
                _line_blocks(builder, [" ".join(el.text.split())], stack)
            for child in el:
                if isinstance(child.tag, str):
                    _walk_html(child, builder, stack, tables)
    if forced and "after" in style:
        builder.end_page()


def parse_html(data: bytes) -> ParsedDocument:
    """Sanitized HTML (scripts, styles and active content removed before text extraction)."""
    root = html.fromstring(data)
    unwanted = cast(
        "list[HtmlElement]", root.xpath("//script|//style|//iframe|//object|//embed|//form")
    )
    for bad in unwanted:
        parent = bad.getparent()
        if parent is not None:
            parent.remove(bad)
    builder = _Builder()
    _walk_html(root, builder, [], [])
    return builder.build("html", "legal")


# ---------------------------------------------------------------- plain text and DOCX
def parse_text(data: bytes) -> ParsedDocument:
    """UTF-8 text; pages are not defined, so the whole text is one page."""
    builder, stack = _Builder(), []
    _line_blocks(
        builder,
        [ln.strip() for ln in data.decode("utf-8", errors="replace").splitlines() if ln.strip()],
        stack,
    )
    return builder.build("txt", "legal")


def parse_docx(data: bytes) -> ParsedDocument:
    """DOCX through python-docx: heading styles become headings, table rows stay together."""
    import docx  # noqa: PLC0415 - optional heavy import

    document = docx.Document(io.BytesIO(data))
    builder, stack, table_no = _Builder(), [], 0
    body = cast("Any", document.element).body  # python-docx types the body loosely
    paragraphs = {p._p: p for p in document.paragraphs}  # noqa: SLF001
    tables = {t._tbl: t for t in document.tables}  # noqa: SLF001
    for child in body.iterchildren():
        if child in paragraphs:
            p = paragraphs[child]
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading"):
                level = int(re.sub(r"\D", "", style) or 1) - 1
                del stack[level:]
                stack.append(p.text.strip())
                builder.add(
                    p.text,
                    "heading",
                    level=level,
                    section_path=" > ".join(stack),
                    heading=p.text.strip(),
                )
            elif p.text.strip():
                _line_blocks(builder, [p.text], stack)
        elif child in tables:
            table_no += 1
            for row in tables[child].rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    builder.add(
                        " | ".join(cells),
                        "table_row",
                        table_id=table_no,
                        section_path=" > ".join(stack),
                    )
    return builder.build("docx", "legal")


# XML parsers: JATS articles, SPL drug labels and eCFR sections
def _xml_text(el: etree._Element) -> str:
    return " ".join(" ".join(str(t) for t in el.itertext()).split())


def _local(tag: object) -> str:
    return etree.QName(tag).localname if isinstance(tag, str) else ""


def parse_xml(data: bytes) -> ParsedDocument:
    """JATS articles, SPL drug labels and eCFR sections (the root element decides)."""
    parser = etree.XMLParser(
        resolve_entities=False, no_network=True, huge_tree=False, recover=False
    )
    root = etree.fromstring(data, parser)
    name = _local(root.tag)
    if name == "article":
        return _parse_jats(root)
    if name == "document":
        return _parse_spl(root)
    return _parse_ecfr(root)


def _table_rows(builder: _Builder, table: etree._Element, number: int, path: str) -> None:
    for row in table.iter():
        if _local(row.tag) == "tr":
            cells = [t for t in (_xml_text(c) for c in row if _local(c.tag) in {"td", "th"}) if t]
            if cells:
                builder.add(" | ".join(cells), "table_row", table_id=number, section_path=path)


def _walk_jats_section(
    builder: _Builder, tables: list[int], sec: etree._Element, stack: list[str]
) -> None:
    heading = sec.find("title")
    label = _xml_text(heading) if heading is not None else ""
    here = [*stack, label] if label else stack
    path = " > ".join(here)
    if label:
        builder.add(label, "heading", level=len(here) - 1, section_path=path, heading=label)
    for child in sec:
        tag = _local(child.tag)
        if tag == "p":
            builder.add(_xml_text(child), "paragraph", section_path=path)
        elif tag == "sec":
            _walk_jats_section(builder, tables, child, here)
        elif tag == "table-wrap":
            tables[0] += 1
            for table in child.iter("table"):
                _table_rows(builder, table, tables[0], path)


def _parse_jats(root: etree._Element) -> ParsedDocument:
    builder, tables = _Builder(), [0]
    title = root.find(".//article-title")
    if title is not None:
        builder.add(
            _xml_text(title),
            "heading",
            level=0,
            section_path=_xml_text(title),
            heading=_xml_text(title),
        )
    abstract = root.find(".//abstract")
    if abstract is not None:
        for p in abstract.iter("p"):
            builder.add(_xml_text(p), "paragraph", section_path="Abstract")
    body = root.find(".//body")
    if body is not None:
        for sec in body.findall("sec"):
            _walk_jats_section(builder, tables, sec, [])
    return builder.build("jats", "medical")


def _parse_spl(root: etree._Element) -> ParsedDocument:
    ns = {"h": "urn:hl7-org:v3"}
    builder, tables = _Builder(), [0]
    title = root.find("h:title", ns)
    if title is not None:
        builder.add(_xml_text(title), "paragraph", section_path="Label")
    for section in root.iterfind(".//h:component/h:section", ns):
        code = section.find("h:code", ns)
        code_value = code.get("code") if code is not None else None
        sec_title = section.find("h:title", ns)
        label = (
            _xml_text(sec_title)
            if sec_title is not None
            else (_SPL_CODES.get(code_value or "", "") or "")
        )
        if not label:
            continue
        path = f"{label} [{code_value}]" if code_value else label
        builder.add(
            label, "heading", level=0, section_path=path, clause_no=code_value, heading=label
        )
        text = section.find("h:text", ns)
        if text is None:
            continue
        for child in text:
            tag = _local(child.tag)
            if tag == "table":
                tables[0] += 1
                _table_rows(builder, child, tables[0], path)
            elif _xml_text(child):
                builder.add(_xml_text(child), "paragraph", section_path=path)
    return builder.build("spl", "medical")


def _parse_ecfr(root: etree._Element) -> ParsedDocument:
    builder = _Builder()
    for div in root.iter():
        if _local(div.tag).startswith("DIV") and div.get("TYPE") == "SECTION":
            number = (div.get("N") or "").strip()
            head = div.find("HEAD")
            title = _xml_text(head) if head is not None else number
            path = title
            builder.add(
                title, "heading", level=0, section_path=path, clause_no=number, heading=title
            )
            for p in div.iter("P"):
                builder.add(_xml_text(p), "paragraph", section_path=path, clause_no=number)
    return builder.build("ecfr", "legal")


def parse(data: bytes, declared: str) -> ParsedDocument:
    """Verify the type, then parse. The caller handles scanning and size limits first."""
    kind = detect_format(data, declared)
    if kind == "pdf":
        return parse_pdf(data)
    if kind == "html":
        return parse_html(data)
    if kind == "xml":
        return parse_xml(data)
    if kind == "docx":
        return parse_docx(data)
    return parse_text(data)
